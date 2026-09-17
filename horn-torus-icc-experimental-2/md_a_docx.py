#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Convierte los documentos de trabajo en markdown a .docx, para pegarlos en la tesis.

Subconjunto soportado: encabezados (#, ##, ###), párrafos, citas en bloque (>),
tablas (| … |), listas numeradas y con guión, reglas (---), imágenes
![pie]({{artifact:art_ID}}) resueltas contra un mapa id -> archivo local, y
énfasis en línea (**negrita**, *cursiva*, `monoespaciado`).

Uso:
    from md_a_docx import md_a_docx
    md_a_docx("capitulo.md", "capitulo.docx", imagenes={"art_xxx": "fig1.png"})
"""

import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

GRIS = RGBColor(0x45, 0x5A, 0x64)


def _inline(par, texto, base_italic=False):
    """Escribe texto con **negrita**, *cursiva* y `mono` en un párrafo."""
    for tok in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)", texto):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            r = par.add_run(tok[2:-2]); r.bold = True
        elif tok.startswith("`") and tok.endswith("`"):
            r = par.add_run(tok[1:-1]); r.font.name = "Consolas"; r.font.size = Pt(9.5)
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            r = par.add_run(tok[1:-1]); r.italic = True
        else:
            r = par.add_run(tok)
        if base_italic:
            r.italic = True


def _tabla(doc, filas):
    cab = [c.strip() for c in filas[0].strip().strip("|").split("|")]
    cuerpo = [[c.strip() for c in f.strip().strip("|").split("|")] for f in filas[2:]]
    ncol = max([len(cab)] + [len(f) for f in cuerpo])
    t = doc.add_table(rows=1, cols=ncol)
    try:
        t.style = "Table Grid"
    except KeyError:
        pass
    for j in range(ncol):
        p = t.rows[0].cells[j].paragraphs[0]
        _inline(p, cab[j] if j < len(cab) else "")
        for r in p.runs:
            r.bold = True
            r.font.size = Pt(9)
    for fila in cuerpo:
        celdas = t.add_row().cells
        for j in range(ncol):
            p = celdas[j].paragraphs[0]
            _inline(p, fila[j] if j < len(fila) else "")
            for r in p.runs:
                r.font.size = Pt(9)
    doc.add_paragraph()


def md_a_docx(entrada, salida, imagenes=None, ancho_img=Cm(15.5)):
    imagenes = imagenes or {}
    lineas = open(entrada, encoding="utf8").read().replace("\r\n", "\n").split("\n")

    doc = Document()
    est = doc.styles["Normal"]
    est.font.name = "Calibri"
    est.font.size = Pt(11)
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2.5)

    i, buf, cita = 0, [], []

    def cerrar_parrafo():
        if buf:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            _inline(p, " ".join(buf))
            buf.clear()

    def cerrar_cita():
        if cita:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.8)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(6)
            _inline(p, " ".join(cita))
            for r in p.runs:
                r.font.size = Pt(10)
            cita.clear()

    while i < len(lineas):
        ln = lineas[i]
        s = ln.strip()

        if s.startswith("|") and i + 1 < len(lineas) and set(lineas[i+1].strip()) <= set("|-: "):
            cerrar_parrafo(); cerrar_cita()
            filas = []
            while i < len(lineas) and lineas[i].strip().startswith("|"):
                filas.append(lineas[i]); i += 1
            _tabla(doc, filas)
            continue

        m = re.match(r"^!\[([^\]]*)\]\(\{\{artifact:([^}]+)\}\}\)", s)
        if m:
            cerrar_parrafo(); cerrar_cita()
            ruta = imagenes.get(m.group(2))
            if ruta:
                doc.add_picture(ruta, width=ancho_img)
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                doc.add_paragraph("[falta la imagen: %s]" % m.group(2))
            i += 1
            continue

        if not s:
            cerrar_parrafo(); cerrar_cita()
        elif s.startswith("> "):
            cerrar_parrafo(); cita.append(s[2:])
        elif s == ">":
            cita.append("")
        elif s.startswith("#"):
            cerrar_parrafo(); cerrar_cita()
            n = len(s) - len(s.lstrip("#"))
            doc.add_heading(re.sub(r"[*`]", "", s[n:].strip()), level=min(n, 4))
        elif s in ("---", "***"):
            cerrar_parrafo(); cerrar_cita()
        elif re.match(r"^(\d+\.|[-*])\s", s):
            cerrar_parrafo(); cerrar_cita()
            estilo = "List Number" if re.match(r"^\d+\.", s) else "List Bullet"
            item = [re.sub(r"^(\d+\.|[-*])\s+", "", s)]
            # absorbe las líneas de continuación (indentadas) del mismo ítem
            while (i + 1 < len(lineas) and lineas[i+1].strip()
                   and lineas[i+1][:1].isspace()
                   and not re.match(r"^\s*(\d+\.|[-*])\s", lineas[i+1])):
                i += 1
                item.append(lineas[i].strip())
            p = doc.add_paragraph(style=estilo)
            _inline(p, " ".join(item))
        elif re.match(r"^\*Figura .*\*$", s):
            cerrar_parrafo(); cerrar_cita()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _inline(p, s.strip("*"), base_italic=True)
            for r in p.runs:
                r.font.size = Pt(9)
                r.font.color.rgb = GRIS
        else:
            cerrar_cita(); buf.append(s)
        i += 1

    cerrar_parrafo(); cerrar_cita()
    doc.save(salida)
    return salida
